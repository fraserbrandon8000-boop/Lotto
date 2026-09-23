"""Step 1: independent audit of the source workbook and construction of the analysis dataset.

Layers kept separate:
  01_original_workbook_extract.csv   rows exactly as found in the workbook
  02_externally_recovered_rows.csv   rows recovered from an official external source (none: see notes)
  03_recent_additions.csv            #2338 and #2339 as supplied in the session prompt
  04_normalized_analysis_dataset.csv chronological union used by every analysis
"""
import csv
import datetime as dt
import os
from collections import Counter

import openpyxl

from common import DATA, NORMALIZED, OUT, WORKBOOK, sha256_file, write_json

HEADER = ["draw_id", "draw_date", "weekday", "n1", "n2", "n3", "n4", "n5", "n6", "bonus", "source", "layer"]

RECENT = [
    (2338, "2026-09-16", [1, 3, 10, 13, 18, 29], 17),
    (2339, "2026-09-19", [1, 4, 6, 13, 23, 28], 20),
]
RECENT_SOURCE = "Supplied by the user in the clean-room session prompt (2026-09-23)"

RECOVERY_NOTE = (
    "Official Supreme Ventures pages (supremeventures.com) were not reachable from this environment: "
    "the network egress proxy denied the connection for both shell and fetch tools. Third-party "
    "aggregators are not an official source, so no rows were recovered and draws 2252-2261 remain a "
    "documented gap. No values were guessed."
)


def main():
    issues = []
    wb = openpyxl.load_workbook(WORKBOOK)
    ws = wb.worksheets[0]
    meta = {
        "workbook_path": os.path.relpath(WORKBOOK, os.path.join(DATA, "..", "..", "..")),
        "workbook_sha256": sha256_file(WORKBOOK),
        "sheets": wb.sheetnames,
        "sheet_dimensions": ws.dimensions,
        "title_rows": [ws.cell(r, 1).value for r in (2, 3, 4)],
        "hidden_rows": [r for r, d in ws.row_dimensions.items() if d.hidden],
        "hidden_columns": [c for c, d in ws.column_dimensions.items() if d.hidden],
        "merged_ranges": [str(m) for m in ws.merged_cells.ranges],
    }
    header = [c.value for c in ws[5]]
    if header != ["Draw No.", "Draw Date", "Number 1", "Number 2", "Number 3", "Number 4", "Number 5", "Number 6", "Bonus", "Source"]:
        issues.append({"check": "header", "detail": f"unexpected header {header}"})

    raw = []
    for r in ws.iter_rows(min_row=6, values_only=True):
        if all(v is None for v in r):
            continue
        raw.append(r)

    # ---- per-row validation --------------------------------------------------------------
    rows = []
    for i, r in enumerate(raw, start=6):
        did, date, *rest = r
        mains, bonus, src = list(rest[:6]), rest[6], rest[7]
        if any(v is None for v in r[:9]):
            issues.append({"check": "missing_values", "row": i, "detail": str(r)})
        if not isinstance(did, int):
            issues.append({"check": "draw_id_type", "row": i, "detail": repr(did)})
        if not isinstance(date, dt.datetime):
            issues.append({"check": "malformed_date", "row": i, "detail": repr(date)})
            continue
        if date.time() != dt.time(0, 0):
            issues.append({"check": "date_has_time_component", "row": i, "detail": str(date)})
        if len(mains) != 6 or not all(isinstance(m, int) for m in mains):
            issues.append({"check": "six_integer_mains", "row": i, "detail": str(mains)})
        if len(set(mains)) != 6:
            issues.append({"check": "unique_mains", "row": i, "detail": str(mains)})
        if not all(1 <= m <= 38 for m in mains):
            issues.append({"check": "main_range", "row": i, "detail": str(mains)})
        if mains != sorted(mains):
            issues.append({"check": "mains_not_sorted_in_source", "row": i, "detail": str(mains)})
        if not (isinstance(bonus, int) and 1 <= bonus <= 38):
            issues.append({"check": "bonus_range", "row": i, "detail": repr(bonus)})
        elif bonus in mains:
            issues.append({"check": "bonus_duplicates_main", "row": i, "detail": f"{mains} bonus {bonus}"})
        wd = date.strftime("%A")
        if wd not in ("Wednesday", "Saturday"):
            issues.append({"check": "weekday", "row": i, "detail": f"{did} {date.date()} {wd}"})
        rows.append({"draw_id": did, "draw_date": date.date().isoformat(), "weekday": wd,
                     "mains": sorted(mains), "bonus": bonus, "source": src, "sheet_row": i})

    # ---- ordering, duplicates, gaps ------------------------------------------------------
    ids = [r["draw_id"] for r in rows]
    order_desc = all(ids[k] > ids[k + 1] for k in range(len(ids) - 1))
    dup_ids = [k for k, v in Counter(ids).items() if v > 1]
    dup_records = [k for k, v in Counter((r["draw_date"], tuple(r["mains"]), r["bonus"]) for r in rows).items() if v > 1]
    dup_combos = [k for k, v in Counter(tuple(r["mains"]) for r in rows).items() if v > 1]
    dup_dates = [k for k, v in Counter(r["draw_date"] for r in rows).items() if v > 1]
    if dup_ids:
        issues.append({"check": "duplicate_ids", "detail": dup_ids})
    if dup_records:
        issues.append({"check": "duplicate_records", "detail": [list(x) for x in dup_records]})
    if dup_dates:
        issues.append({"check": "duplicate_dates", "detail": dup_dates})
    missing = sorted(set(range(min(ids), max(ids) + 1)) - set(ids))

    chrono = sorted(rows, key=lambda r: r["draw_id"])
    date_order_ok = all(chrono[k]["draw_date"] < chrono[k + 1]["draw_date"] for k in range(len(chrono) - 1))
    if not date_order_ok:
        issues.append({"check": "chronology", "detail": "dates not strictly increasing with draw id"})

    # expected number of Wed/Sat draw dates between consecutive observed draws vs id difference
    spacing = []
    for a, b in zip(chrono, chrono[1:]):
        da, db = dt.date.fromisoformat(a["draw_date"]), dt.date.fromisoformat(b["draw_date"])
        slots = sum(1 for k in range(1, (db - da).days + 1)
                    if (da + dt.timedelta(days=k)).weekday() in (2, 5))
        if slots != b["draw_id"] - a["draw_id"]:
            spacing.append({"from": a["draw_id"], "to": b["draw_id"], "from_date": a["draw_date"],
                            "to_date": b["draw_date"], "wed_sat_slots": slots, "id_step": b["draw_id"] - a["draw_id"]})
    gap_dates = []
    a = next(r for r in chrono if r["draw_id"] == 2251) if 2251 in ids else None
    if a:
        d0 = dt.date.fromisoformat(a["draw_date"])
        k = 1
        while len(gap_dates) < len(missing):
            d = d0 + dt.timedelta(days=k)
            if d.weekday() in (2, 5):
                gap_dates.append(d.isoformat())
            k += 1

    # header statements vs data
    title = " ".join(str(t) for t in meta["title_rows"])
    header_claims = {
        "claims_167_draws": "167 draws" in title, "actual_rows": len(rows),
        "claims_newest_first": "Newest first" in title, "actual_newest_first": order_desc,
        "claims_missing_2252_2261": "2252" in title and "2261" in title,
        "actual_missing": missing,
        "claims_last_date_2026-09-12": "12 Sep 2026" in title, "actual_last_date": chrono[-1]["draw_date"],
        "claims_first_date_2025-01-01": "1 Jan 2025" in title, "actual_first_date": chrono[0]["draw_date"],
    }
    source_counts = Counter(r["source"] for r in rows)

    # ---- recent additions: validate and check for conflicts with the workbook -------------
    recent_rows, conflicts = [], []
    for did, date, mains, bonus in RECENT:
        d = dt.date.fromisoformat(date)
        ok = (len(set(mains)) == 6 and all(1 <= m <= 38 for m in mains) and 1 <= bonus <= 38
              and bonus not in mains and d.weekday() in (2, 5))
        if not ok:
            issues.append({"check": "recent_addition_invalid", "detail": did})
        if did in ids:
            w = next(r for r in rows if r["draw_id"] == did)
            if w["mains"] != sorted(mains) or w["bonus"] != bonus or w["draw_date"] != date:
                conflicts.append({"draw_id": did, "workbook": w, "supplied": [date, mains, bonus]})
        recent_rows.append({"draw_id": did, "draw_date": date, "weekday": d.strftime("%A"),
                            "mains": sorted(mains), "bonus": bonus, "source": RECENT_SOURCE})
    last_wb = chrono[-1]
    d_last = dt.date.fromisoformat(last_wb["draw_date"])
    continuity = []
    prev_id, prev_date = last_wb["draw_id"], d_last
    for r in recent_rows:
        d = dt.date.fromisoformat(r["draw_date"])
        slots = sum(1 for k in range(1, (d - prev_date).days + 1) if (prev_date + dt.timedelta(days=k)).weekday() in (2, 5))
        continuity.append({"draw_id": r["draw_id"], "id_step": r["draw_id"] - prev_id, "wed_sat_slots": slots})
        prev_id, prev_date = r["draw_id"], d

    # ---- write layers --------------------------------------------------------------------
    def write(path, recs, layer):
        with open(path, "w", newline="") as f:
            w = csv.writer(f)
            w.writerow(HEADER)
            for r in recs:
                w.writerow([r["draw_id"], r["draw_date"], r["weekday"], *r["mains"], r["bonus"], r["source"], layer])

    write(os.path.join(DATA, "01_original_workbook_extract.csv"), chrono, "original_workbook")
    write(os.path.join(DATA, "02_externally_recovered_rows.csv"), [], "external_official_recovery")
    write(os.path.join(DATA, "03_recent_additions.csv"), recent_rows, "recent_addition_user_supplied")
    final = chrono + [r for r in recent_rows if r["draw_id"] not in ids]
    final.sort(key=lambda r: r["draw_id"])
    for r in final:
        r.setdefault("layer", "original_workbook")
    with open(NORMALIZED, "w", newline="") as f:
        w = csv.writer(f)
        w.writerow(HEADER)
        for r in final:
            layer = "recent_addition_user_supplied" if r["source"] == RECENT_SOURCE else "original_workbook"
            w.writerow([r["draw_id"], r["draw_date"], r["weekday"], *r["mains"], r["bonus"], r["source"], layer])
    with open(os.path.join(DATA, "RECOVERY_NOTES.md"), "w") as f:
        f.write("# Missing draws 2252-2261\n\n" + RECOVERY_NOTE + "\n\nExpected (Wed/Sat schedule) dates of the missing draws: "
                + ", ".join(gap_dates) + "\n")

    ids_final = [r["draw_id"] for r in final]
    report = {
        "workbook": meta,
        "header_claims_vs_data": header_claims,
        "source_column_counts": dict(source_counts),
        "rows_in_workbook": len(rows),
        "draw_id_range_workbook": [min(ids), max(ids)],
        "workbook_order_newest_first": order_desc,
        "missing_ids_in_workbook": missing,
        "missing_ids_expected_dates": dict(zip(missing, gap_dates)),
        "duplicate_ids": dup_ids,
        "duplicate_records": dup_records,
        "duplicate_main_combinations": [list(c) for c in dup_combos],
        "duplicate_dates": dup_dates,
        "chronology_dates_strictly_increasing": date_order_ok,
        "schedule_spacing_anomalies": spacing,
        "recent_additions": recent_rows,
        "recent_additions_conflicts_with_workbook": conflicts,
        "recent_additions_schedule_continuity": continuity,
        "external_recovery": {"rows_recovered": 0, "note": RECOVERY_NOTE},
        "final_dataset": {"rows": len(final), "first_id": ids_final[0], "last_id": ids_final[-1],
                          "first_date": final[0]["draw_date"], "last_date": final[-1]["draw_date"],
                          "missing_ids": sorted(set(range(ids_final[0], ids_final[-1] + 1)) - set(ids_final)),
                          "sha256": sha256_file(NORMALIZED)},
        "issues": issues,
        "issue_count": len(issues),
    }
    write_json(os.path.join(OUT, "01_data_audit.json"), report)
    print({k: report[k] for k in ("rows_in_workbook", "missing_ids_in_workbook", "duplicate_ids", "duplicate_records",
                                  "duplicate_main_combinations", "chronology_dates_strictly_increasing",
                                  "schedule_spacing_anomalies", "recent_additions_conflicts_with_workbook",
                                  "recent_additions_schedule_continuity", "issue_count")})
    print("issues:", issues)
    print("header claims:", header_claims)
    print("final:", report["final_dataset"])


if __name__ == "__main__":
    main()
